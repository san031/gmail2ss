# company name, application status, round#, date of interview, company mail id, link to comapny's website, date applied, role, company review

from pathlib import Path

from openpyxl import Workbook, load_workbook


COLUMNS = [
    "Company", "Status", "Round", "Interview Date", "Company Email",
    "Website", "Date Applied", "Role", "Company Review", "Email IDs",
]
COL = {name: i for i, name in enumerate(COLUMNS)}

STATUSES = (
    "Applied",              # application sent / confirmation received
    "Screening",            # recruiter call, online test, or assignment
    "Interview Scheduled",  # an interview round has a date
    "Interviewed",          # round done, waiting for the result
    "Offer",                # offer received
    "Accepted",             # you accepted the offer
    "Rejected",             # company said no
    "Withdrawn",            # you pulled out
    "Ghosted",              # no reply for a long time
)

DEFAULT_PATH = Path("jobs.xlsx")


def _normalize(text) -> str:
    return " ".join(str(text or "").split()).casefold()


def _open(path: Path):
    if not path.exists():
        wb = Workbook()
        wb.active.title = "Applications"
        wb.active.append(COLUMNS)
        wb.save(path)
    wb = load_workbook(path)
    return wb, wb["Applications"]


def _find_row(ws, company: str, role: str):
    key = (_normalize(company), _normalize(role))
    for row in ws.iter_rows(min_row=2):
        if (_normalize(row[COL["Company"]].value),
                _normalize(row[COL["Role"]].value)) == key:
            return row
    return None


def _get_row_or_fail(ws, company: str, role: str):
    row = _find_row(ws, company, role)
    if row is None:
        raise LookupError(f"No application found for {company} / {role}")
    return row


def _set(row, column: str, value) -> None:
    row[COL[column]].value = value


def _row_to_dict(row) -> dict:
    return {name: row[i].value for name, i in COL.items()}


def _check_status(status: str) -> None:
    if status not in STATUSES:
        raise ValueError(f"Invalid status {status!r}. Allowed: {list(STATUSES)}")


def _email_ids(row) -> list[str]:
    raw = row[COL["Email IDs"]].value or ""
    return [e for e in raw.split(",") if e]


def _already_processed(row, email_id: str) -> bool:
    return bool(email_id) and email_id in _email_ids(row)


def _record_email(row, email_id: str) -> None:
    if email_id:
        _set(row, "Email IDs", ",".join(_email_ids(row) + [email_id]))


def _result(action: str, row) -> dict:
    return {"action": action, "application": _row_to_dict(row)}


def is_email_processed(email_id: str, path: Path = DEFAULT_PATH) -> bool:
    """Check whether any application has already been updated from this email."""
    _, ws = _open(path)
    return any(_already_processed(r, email_id) for r in ws.iter_rows(min_row=2))


def add_application(company: str, role: str, email_id: str = "",
                    date_applied: str = "", company_email: str = "",
                    website: str = "", path: Path = DEFAULT_PATH) -> dict:
    """Track a new application. Does nothing if company + role already exists."""
    wb, ws = _open(path)
    existing = _find_row(ws, company, role)
    if existing:
        return _result("skipped", existing)

    values = {
        "Company": company.strip(), "Status": "Applied", "Round": 0,
        "Interview Date": "", "Company Email": company_email,
        "Website": website, "Date Applied": date_applied,
        "Role": role.strip(), "Company Review": "", "Email IDs": email_id,
    }
    ws.append([values[c] for c in COLUMNS])
    wb.save(path)
    return _result("added", ws[ws.max_row])


def update_status(company: str, role: str, status: str, email_id: str = "",
                  path: Path = DEFAULT_PATH) -> dict:
    """Change the status of an existing application."""
    _check_status(status)
    wb, ws = _open(path)
    row = _get_row_or_fail(ws, company, role)
    # if _already_processed(row, email_id):
    #     return _result("skipped", row)

    _set(row, "Status", status)
    _record_email(row, email_id)
    wb.save(path)
    return _result("updated", row)


def schedule_interview(company: str, role: str, interview_date: str,
                       email_id: str = "", round_number: int | None = None,
                       path: Path = DEFAULT_PATH) -> dict:
    """Record an interview. If round_number is omitted, moves to the next round."""
    wb, ws = _open(path)
    row = _get_row_or_fail(ws, company, role)
    # if _already_processed(row, email_id):
    #     return _result("skipped", row)

    current = row[COL["Round"]].value or 0
    _set(row, "Round", round_number if round_number is not None else current + 1)
    _set(row, "Interview Date", interview_date)
    _set(row, "Status", "Interview Scheduled")
    _record_email(row, email_id)
    wb.save(path)
    return _result("updated", row)


def update_details(company: str, role: str, email_id: str = "",
                   company_email: str | None = None, website: str | None = None,
                   company_review: str | None = None,
                   path: Path = DEFAULT_PATH) -> dict:
    """Fill in or change contact email, website, or your review of the company."""
    wb, ws = _open(path)
    row = _get_row_or_fail(ws, company, role)
    # if _already_processed(row, email_id):
    #     return _result("skipped", row)

    for column, value in [("Company Email", company_email),
                          ("Website", website),
                          ("Company Review", company_review)]:
        if value is not None:
            _set(row, column, value)
    _record_email(row, email_id)
    wb.save(path)
    return _result("updated", row)


def find_application(company: str, role: str,
                     path: Path = DEFAULT_PATH) -> dict | None:
    """Return one application, or None if not tracked."""
    _, ws = _open(path)
    row = _find_row(ws, company, role)
    return _row_to_dict(row) if row else None


def list_applications(status: str | None = None,
                      path: Path = DEFAULT_PATH) -> list[dict]:
    """Return all applications, optionally filtered by status."""
    if status:
        _check_status(status)
    _, ws = _open(path)
    rows = [_row_to_dict(r) for r in ws.iter_rows(min_row=2)]
    return [r for r in rows if not status or r["Status"] == status]
